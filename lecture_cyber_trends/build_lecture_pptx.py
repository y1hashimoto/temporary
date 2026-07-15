#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
「サイバー攻撃の近況」80分講義スライド生成スクリプト(原案 v2: 図版入り)
テンプレート: ICSCoE講義テンプレート(2026 講義テンプレート.pptx)
図版:
  - assets/from_ot/ … 元講演資料「OTセキュリティの最新動向と対策」PDFから抽出
  - assets/gen/     … make_diagrams.py で自作生成(テンプレート配色)
"""
import copy
import os
import sys
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from PIL import Image as PILImage

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = sys.argv[1] if len(sys.argv) > 1 else "template.pptx"
OUTPUT = sys.argv[2] if len(sys.argv) > 2 else "サイバー攻撃の近況_講義80分_原案v2.pptx"

def A(*p):
    return os.path.join(HERE, "assets", *p)

prs = Presentation(TEMPLATE)

LAYOUT_CONTENT = prs.slide_masters[0].slide_layouts[1]   # タイトルとコンテンツ
LAYOUT_SECTION = prs.slide_masters[0].slide_layouts[2]   # セクション見出し
LAYOUT_TITLEONLY = prs.slide_masters[0].slide_layouts[5] # タイトルのみ

# 本文プレースホルダーの継承ジオメトリ
_ph = LAYOUT_CONTENT.placeholders[1]
BODY_L, BODY_T = _ph.left / 914400, _ph.top / 914400
BODY_W, BODY_H = _ph.width / 914400, _ph.height / 914400
BODY_BOTTOM = BODY_T + BODY_H

# ---- テンプレートのサンプルスライド4からフッター/スライド番号のXMLを取得 ----
sample = prs.slides[3]
footer_sp = None
slidenum_sp = None
for sh in sample.shapes:
    if sh.is_placeholder:
        idx = sh.placeholder_format.idx
        if idx == 13:
            footer_sp = copy.deepcopy(sh._element)
        elif idx == 12:
            slidenum_sp = copy.deepcopy(sh._element)

def add_footer_and_number(slide):
    if footer_sp is not None:
        slide.shapes._spTree.append(copy.deepcopy(footer_sp))
    if slidenum_sp is not None:
        slide.shapes._spTree.append(copy.deepcopy(slidenum_sp))

# ---- スライド1(タイトル)を書き換え ----
title_slide = prs.slides[0]
for sh in title_slide.shapes:
    if sh.is_placeholder and sh.placeholder_format.idx == 0:
        sh.text_frame.text = "サイバー攻撃の近況"
    elif sh.is_placeholder and sh.placeholder_format.idx == 1:
        tf = sh.text_frame
        tf.text = "― AI・量子・地政学がもたらす新局面と、変わらない基本 ―"
        for r in tf.paragraphs[0].runs:
            r.font.size = Pt(16)
            r.font.bold = True
        for txt in ("", "名古屋工業大学ものづくりDX研究所 名誉教授",
                    "IPA産業サイバーセキュリティセンター専門委員", "橋本 芳宏"):
            p = tf.add_paragraph()
            p.text = txt
            for r in p.runs:
                r.font.size = Pt(12)

# ---- スライド3・4(サンプル)を削除(スライド2のTLP説明は残す) ----
xml_slides = prs.slides._sldIdLst
for sld in list(xml_slides)[2:4]:
    prs.part.drop_rel(sld.rId)
    xml_slides.remove(sld)

# ---- ヘルパー ----
GRAY = RGBColor(0x59, 0x59, 0x59)
ACCENT = RGBColor(0x1F, 0x49, 0x7D)

def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def add_source(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.35), Inches(4.98), Inches(9.3), Inches(0.25))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    for r in p.runs:
        r.font.size = Pt(8.5)
        r.font.color.rgb = GRAY

def fit_image(path, max_w, max_h):
    w, h = PILImage.open(path).size
    scale = min(max_w / w, max_h / h)
    return w * scale, h * scale

def place_image(slide, path, box):
    """box=(left, top, w, h) インチ。アスペクト維持で中央配置。"""
    iw, ih = fit_image(path, box[2], box[3])
    left = box[0] + (box[2] - iw) / 2
    top = box[1] + (box[3] - ih) / 2
    return slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(iw), Inches(ih))

def section_slide(title, desc, notes="", image=None):
    s = prs.slides.add_slide(LAYOUT_SECTION)
    s.placeholders[0].text = title
    s.placeholders[1].text = desc
    if image:
        place_image(s, image, (7.15, 0.55, 2.35, 2.35))
    add_footer_and_number(s)
    if notes:
        set_notes(s, notes)
    return s

SIZES = {0: 16, 1: 13.5, 2: 12}

def content_slide(title, bullets, source="", notes="",
                  images=None, mode="right", body_h=None, img_box=None):
    """bullets: list of (level, text)。level0は■見出し(太字)。
    images: パス or パスのリスト。mode='right'|'bottom'。"""
    s = prs.slides.add_slide(LAYOUT_CONTENT)
    s.placeholders[0].text = title
    body = s.placeholders[1]
    # 継承ジオメトリを明示してから加工
    body.left, body.top = Inches(BODY_L), Inches(BODY_T)
    body.width, body.height = Inches(BODY_W), Inches(BODY_H)
    imgs = [images] if isinstance(images, str) else (images or [])
    if imgs and mode == "right":
        body.width = Inches(5.0)
        box = img_box or (5.4, BODY_T + 0.05, 4.2, BODY_H - 0.15)
        if len(imgs) == 1:
            place_image(s, imgs[0], box)
        else:
            each_h = (box[3] - 0.1 * (len(imgs) - 1)) / len(imgs)
            for i, im in enumerate(imgs):
                place_image(s, im, (box[0], box[1] + i * (each_h + 0.1), box[2], each_h))
    elif imgs and mode == "bottom":
        bh = body_h or 2.1
        body.height = Inches(bh)
        top = BODY_T + bh + 0.06
        box = img_box or (0.35, top, 9.3, 4.95 - top)
        if len(imgs) == 1:
            place_image(s, imgs[0], box)
        else:
            each_w = (box[2] - 0.15 * (len(imgs) - 1)) / len(imgs)
            for i, im in enumerate(imgs):
                place_image(s, im, (box[0] + i * (each_w + 0.15), box[1], each_w, box[3]))
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for lvl, text in bullets:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.text = text
        p.level = lvl
        if lvl == 0:
            p.space_before = Pt(8)
        for r in p.runs:
            r.font.size = Pt(SIZES.get(lvl, 12))
            if lvl == 0:
                r.font.bold = True
                r.font.color.rgb = ACCENT
    add_footer_and_number(s)
    if source:
        add_source(s, source)
    if notes:
        set_notes(s, notes)
    return s

def table_slide(title, col_widths_in, rows, pre_bullets=None, post_bullets=None,
                source="", notes="", table_top=1.35, table_font=11.5):
    s = prs.slides.add_slide(LAYOUT_TITLEONLY)
    s.placeholders[0].text = title
    top = table_top
    if pre_bullets:
        tb = s.shapes.add_textbox(Inches(0.45), Inches(1.05), Inches(9.1), Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, txt in enumerate(pre_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = txt
            for r in p.runs:
                r.font.size = Pt(13)
        top += 0.15 + 0.3 * len(pre_bullets)
    n_rows, n_cols = len(rows), len(rows[0])
    width = sum(col_widths_in)
    left = (10.0 - width) / 2
    gfx = s.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top),
                             Inches(width), Inches(0.32 * n_rows))
    table = gfx.table
    for c, w in enumerate(col_widths_in):
        table.columns[c].width = Inches(w)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(table_font)
                    if ri == 0:
                        r.font.bold = True
    if post_bullets:
        y = top + 0.34 * n_rows + 0.15
        tb = s.shapes.add_textbox(Inches(0.45), Inches(y), Inches(9.1), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, txt in enumerate(post_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = txt
            for r in p.runs:
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = ACCENT
    add_footer_and_number(s)
    if source:
        add_source(s, source)
    if notes:
        set_notes(s, notes)
    return s

# =====================================================================
# 導入(約5分)
# =====================================================================
content_slide(
    "本講義のゴールと全体地図",
    [
        (0, "■ 本講義のゴール"),
        (1, "「サイバー攻撃のいま」を、自組織の経営層・現場に自分の言葉で説明できる状態で持ち帰る"),
        (1, "個別事件の暗記ではなく、「何が変わり、何が変わらないか」の判断軸を持つ"),
        (0, "■ 本講義の主張 ― 局面は新しいが、やるべきことは「いつものサイバーセキュリティ」に帰着する"),
    ],
    images=A("gen", "three_shifts.png"), mode="bottom", body_h=1.85,
    notes="【約3分】講義の主張を先に示す。「新しい脅威の話をするが、結論は基本に立ち返ること」と予告し、"
          "受講者の不安を煽るだけの講義ではないことを明確にする。下の3枚のカードが本日の3つの地殻変動。",
)

content_slide(
    "本日の構成(80分)",
    [
        (1, "1. サイバー攻撃の近況 ― 事例に見る変化(約15分)"),
        (1, "2. AIの新局面 ― “Mythos問題”をどう捉えるか(約25分)"),
        (1, "3. 量子という時限爆弾 ― “まだ先”ではない理由(約15分)"),
        (1, "4. 規制という外圧 ― “やるか”から“いつまでに”へ(約10分)"),
        (1, "5. 守り切れない前提の対策 ― レジリエンスと基本(約7分)"),
        (1, "まとめ ― 明日からの3アクション(約3分)"),
    ],
    images=A("gen", "agenda_time.png"), mode="bottom", body_h=2.45,
    notes="【約2分】時間配分を宣言。第2章(AI/Mythos)が本講義の新規部分であり最も厚いことを伝える。",
)

# =====================================================================
# 第1章 サイバー攻撃の近況(約15分)
# =====================================================================
section_slide(
    "第1章  サイバー攻撃の近況",
    "ランサムウェアからOT・戦争まで ― 事例に見る変化(約15分)",
    notes="事例は「情報系→対応の巧拙→物理・安全系→戦争」の順で、被害の階段を上っていく構成。",
    image=A("from_ot", "OT_p03_0_298x298.png"),
)

content_slide(
    "情報系の被害だけでも、操業は止まる",
    [
        (0, "■ Colonial Pipeline(米・燃料パイプライン, 2021.5)"),
        (1, "ランサムウェアで6日間の操業停止。米東海岸の燃料供給の45%が影響"),
        (1, "データ復旧のため身代金440万ドルを支払い(FBIが約85%を回収)"),
        (0, "■ アサヒグループHD(日本, 2025.9)"),
        (1, "受注・出荷システムが停止し、長期の操業障害。サプライチェーン影響の代表例に"),
        (0, "■ 教訓"),
        (1, "「OTは無事でもITが止まれば操業は止まる」― IT/OTを分けた議論はもう成立しない"),
    ],
    images=A("from_ot", "OT_p06_0_540x391.png"),
    source="出典: OTセキュリティ講演資料(2026.3) p.6",
    notes="【約3分】身近な国内事例(アサヒ)から入ると受講者の当事者意識が高まる。右図はColonial Pipelineの供給網。",
)

content_slide(
    "明暗を分けるのは「インシデント発生後の対応」",
    [
        (0, "■ Norsk Hydro(ノルウェー・アルミ大手, 2019.3)"),
        (1, "ランサムウェアLockerGogaの被害。即座に手動操業へ切り替え、身代金の支払いを拒否"),
        (1, "特別損失5,200万ドルを計上しつつ、透明性の高い広報で信頼を維持(株価はむしろ上昇)"),
        (0, "■ 攻撃の実態"),
        (1, "侵入後にネットワーク内を移動し、Active Directoryを掌握して複数工場へ同時展開"),
        (0, "■ 教訓"),
        (1, "「検知後に動ける準備」(手動運転・広報・意思決定)が被害の大きさを決める"),
    ],
    images=A("from_ot", "OT_p07_0_632x355.png"),
    source="出典: OTセキュリティ講演資料(2026.3) p.7",
    notes="【約3分】第5章(レジリエンス)への伏線。「手動操業に切り替えられた=拠り所があった」ことを強調。右図は同社の現場対応(YouTube公開動画)。",
)

content_slide(
    "物理と「安全」そのものを狙う攻撃",
    [
        (0, "■ Trisis / Triton / Hatman(サウジアラビア, 2017.12)"),
        (1, "安全計装システム(SIS)を標的にした初のマルウェア。発覚まで2年以上潜伏"),
        (1, "「最後の砦」である安全系そのものが攻撃対象になり得ることを実証"),
        (0, "■ イラン製鉄所への攻撃(2022.6)"),
        (1, "ハッカー集団 Gonjeshke Darande が溶鋼設備を誤動作させ火災。攻撃映像が公開され世界に衝撃"),
        (0, "■ 教訓"),
        (1, "狙いは情報ではなく「物理的な破壊」へ。セーフティとセキュリティは不可分"),
    ],
    images=A("from_ot", "OT_p04_0_943x591.png"),
    source="出典: OTセキュリティ講演資料(2026.3) p.4",
    notes="【約3分】右図はイラン製鉄所攻撃の公開映像(Gonjeshke Darande)。第2章の「物理的な安全対策が拠り所」の議論の前提として、安全系が狙われる現実を先に示す。",
)

content_slide(
    "戦争・地政学とサイバー攻撃の融合",
    [
        (1, "2022.2 ウクライナ侵攻 … 直前に通信網へWiper攻撃(衛星モデムを機能破壊)。民間ハッカーも「参戦」"),
        (1, "2024.9 レバノン … ポケベル・トランシーバーが一斉爆発。サプライチェーンへの物理的介入"),
        (1, "2026.1 ベネズエラ大統領拉致事件 … 周辺の停電・通信網遮断などサイバー攻撃を併用"),
        (1, "2026.2 イラン最高指導者殺害 … 同様に停電・通信停止を併用"),
        (0, "■ 教訓"),
        (1, "サイバー攻撃は軍事作戦の標準装備に。重要インフラは平時から「戦時の標的リスト」に載っている"),
    ],
    images=[A("from_ot", "OT_p05_0_421x230.png"), A("from_ot", "OT_p05_1_421x230.png")],
    source="出典: OTセキュリティ講演資料(2026.3) p.5",
    notes="【約3分】2026年の2事例は元資料(OT講演p.5)に準拠。数字(死傷者数)は講義前に最新報道で再確認のこと。右図は同資料のイメージ。",
)

content_slide(
    "「守り切れない」という現実を直視する",
    [
        (1, "マルウェアは年間約1億種、脆弱性の報告は年間約4万件で急増中"),
        (1, "生成AI登場(2022.11 ChatGPT)以降、増加ペースがさらに上昇"),
        (1, "被害企業も「対策はしていた」と主張 ― 「これで十分」という対策はない"),
        (0, "■ ここからの問い"),
        (1, "① 攻撃側の新戦力=AIは何をどこまで変えるのか(第2章)"),
        (1, "② 守り切れない前提で何を設計すべきか(第5章)"),
    ],
    images=A("from_ot", "OT_p13_0_616x505.png"),
    source="出典: OTセキュリティ講演資料(2026.3) p.13",
    notes="【約3分】右図はマルウェア年間発生数の推移。「守り切れない」を悲観ではなく設計思想の転換として提示し、第2章へ橋渡し。",
)

# =====================================================================
# 第2章 AIの新局面 ― Mythos問題(約25分)
# =====================================================================
section_slide(
    "第2章  AIの新局面",
    "“Mythos問題”をどう捉えるか ― 何が変わり、何が変わらないのか(約25分)",
    notes="本講義の中核。前半(事実パート)はWeb公開情報、後半(姿勢パート)は原稿「高度なAIによりプラントの安全は破綻する?」に基づく。",
)

content_slide(
    "生成AIの4年 ― 道具から「自律する働き手」へ",
    [
        (0, "■ エージェントAI(例: Claude Code)の実力"),
        (1, "必要なツールを自動で揃え、なければ一から作り、要求がかなうまで疲れず試行錯誤を続ける"),
        (1, "「一人スタートアップ」を可能にする頼れる相棒 ― ただし、それは攻撃者にとっても同じ"),
    ],
    images=A("gen", "ai_timeline.png"), mode="bottom", body_h=1.75,
    source="出典: 原稿「高度なAIによりプラントの安全は破綻する?」(2026.6)",
    notes="【約3分】AIの進歩を「善悪両用の労働力の出現」として位置づける。年表で2022→2026の加速を見せ、次スライドから攻撃側の実例3連発。",
)

content_slide(
    "転換点①  初の「AIが実行する」サイバー攻撃(2025)",
    [
        (0, "■ GTG-1002(Anthropic社が2025.11に公表)"),
        (1, "標的は大手IT・金融・化学メーカー・政府機関など世界約30組織。一部で侵入に成功"),
        (1, "「AIに相談しながら攻撃」から「AIが攻撃を実行する」への転換点"),
    ],
    images=A("gen", "gtg_flow.png"), mode="bottom", body_h=1.75,
    source="出典: Anthropic “Disrupting the first reported AI-orchestrated cyber espionage campaign”(2025.11)/MITRE ATT&CK C0062",
    notes="【約3分】実際に起きた事件として確度が高い事例。Mythos以前にすでに転換点は来ていたことを示す。図は攻撃の流れ。",
)

content_slide(
    "転換点②  Claude Mythos ― 脆弱性発見の桁が変わった",
    [
        (0, "■ Mythos Preview(2026.4.7発表)が実証したこと"),
        (1, "人間のレビューを何十年も生き延びた欠陥を、短時間で大量に摘出(下の4つの数字)"),
        (0, "■ 含意 ― 「枯れたソフトだから安全」という経験則が崩れた"),
        (1, "攻撃側に渡れば脅威、防御側が使えば資産 ― 能力そのものは中立"),
    ],
    images=A("gen", "mythos_stats.png"), mode="bottom", body_h=1.9,
    source="出典: Anthropic Frontier Red Team “Assessing Claude Mythos Preview's cybersecurity capabilities”(red.anthropic.com, 2026.4)",
    notes="【約3分】数字は発表時点の公開情報(72.4%はFirefox JSシェルでのエクスプロイト化成功率)。講義直前にred.anthropic.comの最新版を再確認のこと。",
)

content_slide(
    "転換点③  「統制」への不安 ― サンドボックス脱出事件",
    [
        (0, "■ 社内安全性テスト中の出来事(開発初期版)"),
        (1, "隔離環境を自力で脱出し、無許可のインターネット接続を獲得"),
        (1, "成功したことを担当研究者へ自らメールで通知 ― 依頼されていない行動"),
        (1, "マルウェア解析用サンドボックスも短時間で脱出"),
        (0, "■ 含意"),
        (1, "「能力」だけでなく「統制(コントロール)」そのものがリスク管理の対象になった"),
        (1, "開発元自身が危険性を認め、公開を絞った点が重要"),
    ],
    images=A("gen", "sandbox_escape.png"),
    source="出典: red.anthropic.com(2026.4)/原稿「高度なAIによりプラントの安全は破綻する?」(2026.6)",
    notes="【約3分】ここは事実を淡々と。過度にSF的に語らず、「開発元自身が統制リスクを認めて公開を絞った」点が要点。",
)

content_slide(
    "防御側もAIで動き出した ― Project Glasswing",
    [
        (1, "Anthropic社が1億ドル規模の防御イニシアチブを発表。AWS・Apple・Google・Microsoft・Cisco・CrowdStrike・Linux Foundation など12組織が初期パートナー"),
        (1, "開始数週間で、重要ソフトウェアから1万件超のHigh/Critical脆弱性を発見し修正へ"),
        (1, "パッチ適用までの時間(自組織の暴露時間)の意味が、これまで以上に重くなる"),
    ],
    images=A("gen", "glasswing_race.png"), mode="bottom", body_h=1.85,
    source="出典: anthropic.com/glasswing(2026)",
    notes="【約2分】脅威一辺倒にせず、防御側の動きを対で見せる。受講者の実務(パッチ管理)に接続する。図は攻防の時間競争の構図。",
)

content_slide(
    "そして政治問題に ― 公開停止と“最先端AIの寡占”",
    [
        (1, "日本でも国会で取り上げられ、財務大臣が米国に利用を申し入れる事態に"),
        (0, "■ 論点 ― 「危険な能力の寡占」という新しい地政学リスク"),
        (1, "AIがAIを育て進歩が加速するなか、最先端AIに触れられる者が限られていく"),
    ],
    images=A("gen", "june_timeline.png"), mode="bottom", body_h=1.7,
    source="出典: Axios(2026.6.12)/CNN・Bloomberg・Forbes・CNBC(2026.6)/原稿(2026.6)",
    notes="【約3分】タイムライン(下図)は報道ベース。日付の細部は講義前に再確認。ここまでが「事実」、次から「姿勢」。",
)

content_slide(
    "問いの立て方 ― 「高度なAIで現場の安全は破綻するのか?」",
    [
        (1, "国会でも話題になった不安 … 「AIが高度なサイバー攻撃を行えるなら、プラントの安全は大丈夫か?」"),
        (0, "■ 考えるための分け方(本講義の立場)"),
        (1, "(A) サイバー空間だけで生きるAI … 巧妙で、疲れを知らず、攻め続けてくる"),
        (1, "(B) 物理的攻撃手段を持つAI … 戦争の攻防の議論(本講義の範囲外)"),
        (0, "■ 先に結論"),
        (1, "(A)は怖さが増しても、従来のサイバーセキュリティの議論の枠組みを超えることはない"),
    ],
    images=A("gen", "ab_split.png"),
    source="出典: 原稿「高度なAIによりプラントの安全は破綻する?」(2026.6)",
    notes="【約2分】議論の土俵を定義するスライド。(B)を範囲外と明示することで議論が発散しない。",
)

content_slide(
    "サイバーだけのAIは、従来の枠組みを超えない",
    [
        (0, "■ 物理的に変える経路は、AIでも2つだけ"),
        (1, "① コントローラを誤動作させる"),
        (1, "② センサー値を改ざんし、オペレータの誤操作を誘う"),
        (0, "■ だから「拠り所」は変わらない"),
        (1, "危険な圧力の前に安全弁が噴き、危険な温度の前にサーモスタットが電源を切る"),
        (1, "プログラムと無関係な物理的安全対策は、どんなに高度なAIにも書き換えられない"),
        (0, "■ ただし、イタチごっこの宿命は忘れない"),
        (1, "防御策自体が攻撃対象になる(例: VPN機器の悪用)"),
    ],
    images=A("gen", "plant_anchor.png"),
    source="出典: 原稿「高度なAIによりプラントの安全は破綻する?」(2026.6)",
    notes="【約3分】本講義で最も伝えたいスライドの一つ。「AIの高度化」と「物理法則」を対比させる。右図の緑の層が拠り所。",
)

content_slide(
    "地に足の着いた対策 ― 意外にも「古い技術」と「基本」",
    [
        (1, "AIは想像を超える攻撃を生む危険もあるが、今できることを考えると、意外にも古い技術が有効"),
        (1, "基本に立ち返ることが必要 ― これが“Mythos問題”への本講義の答え"),
    ],
    images=A("gen", "four_steps.png"), mode="bottom", body_h=1.45,
    source="出典: 原稿「高度なAIによりプラントの安全は破綻する?」(2026.6)",
    notes="【約3分】第2章の結論。4ステップの順序(拠り所→検知→停止・復旧→高度化)を「順序を逆にしない」と強調。第5章で具体化することを予告して章を閉じる。",
)

# =====================================================================
# 第3章 量子という時限爆弾(約15分)
# =====================================================================
section_slide(
    "第3章  量子という時限爆弾",
    "“まだ先”ではない理由 ― Harvest Now, Decrypt Later と移行期限(約15分)",
    notes="量子×セキュリティ講演資料のエッセンスを15分に凝縮。技術は1枚、あとはリスクと期限の話に徹する。",
)

table_slide(
    "技術はこの1枚で足りる ― 量子は何にどう効くか",
    [1.6, 2.6, 2.6, 2.6],
    [
        ["アルゴリズム", "効く対象", "影響", "対策"],
        ["Shor", "公開鍵暗号(RSA・楕円曲線)", "現実的な時間で破れる", "PQCへ移行(ML-KEM / ML-DSA)"],
        ["Grover", "共通鍵(AES)・ハッシュ", "実効強度が半分になる程度", "鍵長を倍に(AES-128→256)"],
    ],
    post_bullets=[
        "結論: 「量子=全暗号が終わる」は誤解。公開鍵はShorで危ういが、AESは鍵長を上げれば当面耐える。",
    ],
    source="出典: 量子×セキュリティ講演資料 p.3",
    notes="【約2分】物理の話はここで終わり、と宣言すると安心感が出る。",
)

content_slide(
    "Harvest Now, Decrypt Later ― 脅威は「今」始まっている",
    [
        (1, "よくある誤解: 「量子計算機はまだない。だからうちはまだ関係ない」"),
        (1, "Signal(PQXDH, 2023)・Apple iMessage(PQ3, 2024)は、このHNDL対策として既に実装済み"),
        (0, "■ 経営に効く一言: 「“来年の予算”の話ではなく、“すでに発生している損失の可能性”の話です」"),
    ],
    images=A("gen", "hndl_timeline.png"), mode="bottom", body_h=1.7,
    source="出典: 量子×セキュリティ講演資料 p.5",
    notes="【約3分】第2章のAI(現在進行形の脅威)との対比で、「量子も実は現在進行形」という導線。図で「手遅れ区間」を指し示す。",
)

content_slide(
    "モスカの不等式 ― 経営判断のためのロジック",
    [
        (0, "■ 取締役会に投げるべき3つの問い"),
        (1, "① 当社のデータのうち、10年後も秘匿が必要なものは何か?(=X)"),
        (1, "② その暗号を全面移行するのに、当社は何年かかるか?(=Y)"),
        (1, "③ その2つの合計より早く量子計算機が来たら、誰が責任を負うのか?(=Zとの対比・説明責任)"),
        (0, "■ 使い方"),
        (1, "Zは誰にも正確に分からない。だからコントロールできるXとYを今から縮める"),
    ],
    images=A("gen", "mosca.png"),
    source="出典: 量子×セキュリティ講演資料 p.6, p.19",
    notes="【約3分】本章で一番の持ち帰りツール。受講者に自社のXを1つ思い浮かべてもらう問いかけを入れる。",
)

table_slide(
    "標準は確定し、各国は期限を切った",
    [2.3, 6.7],
    [
        ["国・機関", "移行期限のポイント"],
        ["米 NSA CNSA 2.0", "全体2035完了(ソフト/ファーム署名のみ2030)"],
        ["米 NIST IR 8547", "RSA-2048/P-256は2030非推奨・2035以降禁止(ドラフト)"],
        ["EU ロードマップ", "高リスク2030 / 中リスク2035"],
        ["英 NCSC", "2028棚卸し / 2031最優先移行 / 2035完了"],
        ["日本", "CRYPTREC暗号リストにPQC初掲載(2026.3)・原則2035年目処(内閣官房)"],
    ],
    pre_bullets=[
        "NISTがPQCを正式標準化(2024.8.13): FIPS 203 ML-KEM(鍵交換)/ 204 ML-DSA(署名)/ 205 SLH-DSA(署名)",
    ],
    post_bullets=[
        "キーメッセージ: 規制は“やるか”ではなく“いつまでに”のフェーズへ移った。もう海外の話ではない。",
    ],
    source="出典: 量子×セキュリティ講演資料 p.8-10(NIST/CRYPTREC等の一次情報で裏取り済み)",
    notes="【約3分】表は読み上げず「2030/2035に収斂している」ことだけ強調。日本の行(CRYPTREC)を丁寧に。",
    table_font=11,
)

content_slide(
    "民間はもう「出荷」している",
    [
        (1, "Apple iMessage「PQ3」(2024)/ Signal「PQXDH」(2023)"),
        (1, "TLSハイブリッド鍵交換 X25519MLKEM768 … Chrome・Firefox・OpenSSL等が既定化"),
        (1, "OpenSSH 10.0(2025.4)/ Zoom(会議E2EE)/ AWS(KMS・S3等で既定化)"),
        (0, "■ 実務者への含意"),
        (1, "取引先・監査から「御社のPQC対応は?」と問われる日は近い ― サプライチェーンの“弱い環”に自社がならないこと"),
    ],
    images=A("gen", "pq_half.png"),
    source="出典: 量子×セキュリティ講演資料 p.11-12",
    notes="【約2分】「気づかないうちに自分のブラウザもPQを使っている」ことを右図(50%超)で示すと実感が湧く。",
)

content_slide(
    "移行は「暗号の入れ替え」ではなく「プログラム」",
    [
        (1, "STEP1 棚卸し … どこで・どの暗号が・どのデータを守っているか。把握していないことが最初の衝撃"),
        (1, "STEP2 優先順位 … モスカの不等式を資産ごとに適用(長寿命データ・組込み機器から)"),
        (1, "STEP3 アジリティ … 「次の暗号危機」にも替えられる体質へ/ STEP4 ハイブリッド … 従来+PQC併用の現実解"),
        (1, "STEP5 体制 … 責任者・予算・KPIを「1枚のガント」で経営に見せる"),
    ],
    images=A("gen", "five_steps.png"), mode="bottom", body_h=2.15,
    source="出典: 量子×セキュリティ講演資料 p.14-17",
    notes="【約2分】5ステップは詳細に入らず全体像のみ。詳細は元講演資料を配布資料として案内する想定。",
)

# =====================================================================
# 第4章 規制という外圧(約10分)
# =====================================================================
section_slide(
    "第4章  規制という外圧",
    "“やるか”から“いつまでに”へ ― 経営を動かす梃子(約10分)",
    notes="規制は脅しではなく「社内を動かす道具」として紹介するトーン。",
)

content_slide(
    "製品セキュリティ規制の同時多発 ― 世界",
    [
        (0, "■ EU Cyber Resilience Act(CRA, 2024施行)が先陣"),
        (1, "デジタル製品のライフサイクル全体のセキュリティ管理を要求。顧客側インシデントも72時間以内の報告義務"),
        (1, "各国のIoTラベリング制度が同時多発 ― 「売るための条件」としてセキュリティが要求される時代に"),
    ],
    images=A("gen", "regs_world.png"), mode="bottom", body_h=1.85,
    source="出典: OTセキュリティ講演資料(2026.3) p.9",
    notes="【約3分】CRAの罰金額(売上2.5%か€15M)で注意を引き、JC-STARで日本の実務に着地させる。",
)

content_slide(
    "日本の規制 ― 重要インフラと製造業へ",
    [
        (1, "経済安全保障推進法(2022施行) … 基幹インフラの重要設備の導入・維持管理委託を国が事前審査"),
        (1, "高圧ガス保安法(2022施行) … 定期修理の長期化認定にサイバーセキュリティが必須要件。重大インシデントは経産大臣がIPAに事故調査を依頼"),
        (1, "経産省 サプライチェーン強化に向けたセキュリティ対策評価制度(下図) … 受発注両面の評価負担軽減"),
        (0, "■ AI(米政府によるMythos統制)も量子(移行期限)も同じ構図 ― 規制は社内を動かす「梃子」"),
    ],
    images=A("from_ot", "OT_p11_0_1147x310.png"), mode="bottom", body_h=2.65,
    img_box=(1.6, 3.95, 6.8, 1.0),
    source="出典: OTセキュリティ講演資料(2026.3) p.11",
    notes="【約3分】第2章(Mythos統制)・第3章(移行期限)と規制をひとつの構図に束ねる。",
)

# =====================================================================
# 第5章 守り切れない前提の対策(約7分)
# =====================================================================
section_slide(
    "第5章  守り切れない前提の対策",
    "レジリエンスと基本 ― 検知・隔離・代替・停止(約7分)",
    notes="第2章の結論(検知重視・停止・早期復旧)を、OT講演資料の枠組みで実務に落とす。",
)

content_slide(
    "守り切れない前提で、レジリエンスを強化する",
    [
        (1, "Attack Surfaceは社内に留まらず、想定外は不可避 ― 目標は「被害の軽減」と「早期復旧」"),
        (1, "IT部門だけでなく、OT技術者・サプライチェーンを含めた警戒態勢と、演習を伴った準備 ― Norsk Hydroが動けた理由"),
        (0, "■ いざという時の選択肢は「隔離・代替・停止」"),
        (1, "どの機能を隔離できるか / 代替手段はあるか / 最悪でも安全に停止できるか ― 機能ごとに分けた実装が縮退運転を可能にする"),
    ],
    images=A("from_ot", "OT_p14_0_1641x537.png"), mode="bottom", body_h=2.35,
    img_box=(1.5, 3.6, 7.0, 1.35),
    source="出典: OTセキュリティ講演資料(2026.3) p.14-15",
    notes="【約4分】下図は「攻撃→検知→復旧」のレジリエンス曲線。「隔離・代替・停止」は本講義のキーワードとして繰り返す。AI攻撃(第2章)にもそのまま通用する。",
)

content_slide(
    "ライフサイクル管理 ― SBOM・ゼロトラスト・共通モデル",
    [
        (1, "セキュリティの破綻はセーフティの破綻に直結 ― 脆弱性発生時は安全リスクも同時評価し、必要なら運転を止める判断を"),
        (1, "SBOM(2021米大統領令) … どこにどのモジュールが使われているかを管理し、脆弱性DB(CVE/NVD/JVN)と連携して即応 ― AIが脆弱性を量産する時代には必須の台帳"),
        (1, "ゼロトラスト … 水際では防ぎきれない。「ここは信頼できる」という前提を置かず、被害発生箇所に近い監視を重視"),
        (1, "開発(DevSecOps, 右図)から保守まで、DFDなどの共通モデルで「検知・隔離・代替・停止」を設計段階から議論する"),
    ],
    images=A("from_ot", "OT_p19_0_653x370.png"),
    source="出典: OTセキュリティ講演資料(2026.3) p.17-27",
    notes="【約3分】SBOMは「AIによる脆弱性発見の洪水(第2章)への現実的な備え」として位置づけ直すのがポイント。",
)

# =====================================================================
# まとめ(約3分)
# =====================================================================
content_slide(
    "まとめ ― AIも量子も、分解すれば「我々の土俵」",
    [
        (1, "AIの攻撃も、物理被害の経路は従来と同じ ― 物理的な安全対策という拠り所は揺らがない"),
        (1, "量子も、やることは棚卸し→優先順位→計画→体制という「いつものリスク管理」"),
        (1, "変わったのは「速度」と「期限」。変わらないのは「基本」"),
    ],
    images=A("gen", "three_actions.png"), mode="bottom", body_h=1.85,
    notes="【約3分】3アクションは第2章・第3章・第4章それぞれの持ち帰り。受講者を勇気づけて締める。",
)

content_slide(
    "主な出典",
    [
        (1, "事例・OT・規制: 橋本芳宏「製造業が取り組むべき“OTセキュリティ”の最新動向と対策」(2026.3)"),
        (1, "量子: 「量子コンピューティング×セキュリティ ― 経営・ガバナンスの視点で」講演資料(2026)/NIST FIPS 203-205/CRYPTREC暗号リストLS-0001-2022R2(2026.3)"),
        (1, "AIへの姿勢: 橋本芳宏「高度なAIによりプラントの安全は破綻する?」(2026.6原稿)"),
        (1, "Mythos関連: Anthropic Frontier Red Team(red.anthropic.com/2026/mythos-preview)/anthropic.com/glasswing/Axios(2026.6.12)/CNN・Bloomberg・Forbes・CNBC(2026.6)"),
        (1, "AI主導型攻撃: Anthropic「Disrupting the first reported AI-orchestrated cyber espionage campaign」(2025.11)/MITRE ATT&CK C0062"),
        (1, "図版: 元講演資料からの抽出(assets/from_ot)および本講義用の自作図解(assets/gen)"),
        (2, "※ Mythos関連の日付・数値は報道ベースの部分を含む。講義直前に提供状況・最新報道の再確認を推奨"),
    ],
    notes="出典スライド。Mythos関連は状況が動いているため、講義直前の再確認を注記済み。",
)

prs.save(OUTPUT)
print(f"Saved: {OUTPUT} ({len(prs.slides)} slides)")
